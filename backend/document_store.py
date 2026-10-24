from __future__ import annotations

import csv
import json
import os
import re
import sqlite3
import uuid
from collections import Counter
from pathlib import Path
from typing import Any

import fitz
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

ALLOWED_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".html", ".htm",
    ".pdf", ".docx", ".xlsx", ".pptx",
    ".png", ".jpg", ".jpeg",
}
MAX_UPLOAD_BYTES = 15 * 1024 * 1024

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = DATA_DIR / "sovereignai.db"


def init_db() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS documents (
            id TEXT PRIMARY KEY,
            original_name TEXT NOT NULL,
            stored_name TEXT NOT NULL,
            extension TEXT NOT NULL,
            size_bytes INTEGER NOT NULL,
            uploaded_at TEXT NOT NULL,
            file_path TEXT NOT NULL,
            status TEXT NOT NULL,
            metadata TEXT NOT NULL DEFAULT '{}'
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS document_chunks (
            id TEXT PRIMARY KEY,
            document_id TEXT NOT NULL,
            chunk_index INTEGER NOT NULL,
            content TEXT NOT NULL,
            page_number INTEGER,
            token_count INTEGER NOT NULL,
            score REAL NOT NULL DEFAULT 0.0,
            FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
        )
        """
    )
    conn.commit()
    conn.close()


def sanitize_filename(filename: str) -> str:
    candidate = Path(filename).name
    if not candidate or candidate in {".", ".."}:
        raise ValueError("Invalid filename.")
    if os.path.basename(candidate) != candidate:
        raise ValueError("Unsupported filename path.")
    return candidate


def allowed_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def validate_upload(filename: str, size_bytes: int) -> None:
    if not allowed_file(filename):
        raise ValueError(f"Unsupported file type: {Path(filename).suffix or 'unknown'}")
    if size_bytes <= 0:
        raise ValueError("Uploaded file is empty.")
    if size_bytes > MAX_UPLOAD_BYTES:
        raise ValueError(f"File size exceeds {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.")


def _extract_pdf_text(path: Path) -> tuple[str, list[dict[str, Any]]]:
    pages: list[dict[str, Any]] = []
    text_parts: list[str] = []
    document = fitz.open(path)
    for page_number, page in enumerate(document, start=1):
        page_text = page.get_text("text")
        text_parts.append(page_text.strip())
        pages.append({"page_number": page_number, "text": page_text.strip()})
    document.close()
    return "\n\n".join(text_parts).strip(), pages


def _extract_docx_text(path: Path) -> str:
    document = DocxDocument(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs)


def _extract_xlsx_text(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cleaned = [str(value) if value is not None else "" for value in row]
            if cleaned and any(cell.strip() for cell in cleaned):
                rows.append(" | ".join(cleaned))
    return "\n".join(rows)


def _extract_pptx_text(path: Path) -> str:
    presentation = Presentation(path)
    blocks: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            blocks.append(f"Slide {slide_number}: {' '.join(texts)}")
    return "\n".join(blocks)


def _extract_image_text(path: Path) -> str:
    try:
        with Image.open(path) as image:
            width, height = image.size
            metadata = image.getexif()
            info = f"Image file: {path.name}; dimensions: {width}x{height}; format: {image.format or 'unknown'}"
            if metadata:
                info += "; exif_keys=" + ", ".join(str(key) for key in metadata.keys())
            return info
    except Exception:
        return f"Image file: {path.name}; OCR is not enabled in this environment."


def extract_document_text(path: Path) -> tuple[str, list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".json", ".html", ".htm", ".csv"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return text, [{"page_number": 1, "text": text}]
    if suffix == ".pdf":
        return _extract_pdf_text(path)
    if suffix == ".docx":
        return _extract_docx_text(path), [{"page_number": 1, "text": _extract_docx_text(path)}]
    if suffix == ".xlsx":
        return _extract_xlsx_text(path), [{"page_number": 1, "text": _extract_xlsx_text(path)}]
    if suffix == ".pptx":
        return _extract_pptx_text(path), [{"page_number": 1, "text": _extract_pptx_text(path)}]
    if suffix in {".png", ".jpg", ".jpeg"}:
        return _extract_image_text(path), [{"page_number": 1, "text": _extract_image_text(path)}]
    raise ValueError(f"Unsupported document type for extraction: {suffix}")


def chunk_text(text: str, chunk_size: int = 600, overlap: int = 120) -> list[str]:
    if not text.strip():
        return []
    words = text.split()
    if not words:
        return []
    chunks: list[str] = []
    stride = max(chunk_size - overlap, 1)
    for start in range(0, len(words), stride):
        piece = words[start:start + chunk_size]
        if not piece:
            break
        chunks.append(" ".join(piece))
        if len(piece) < chunk_size:
            break
    return chunks


def _tokenize(value: str) -> list[str]:
    return [token.lower() for token in re.findall(r"[A-Za-z0-9]+", value)]


def _lexical_similarity(query: str, text: str) -> float:
    query_tokens = Counter(_tokenize(query))
    text_tokens = Counter(_tokenize(text))
    if not query_tokens or not text_tokens:
        return 0.0
    overlap = sum(min(query_tokens[token], text_tokens[token]) for token in query_tokens)
    return overlap / max(len(query_tokens), 1)


def list_documents() -> list[dict[str, Any]]:
    init_db()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        "SELECT id, original_name, stored_name, extension, size_bytes, uploaded_at, file_path, status, metadata FROM documents ORDER BY uploaded_at DESC"
    ).fetchall()
    conn.close()
    return [dict(row) for row in rows]


def upload_document(file_name: str, file_bytes: bytes) -> dict[str, Any]:
    validate_upload(file_name, len(file_bytes))
    init_db()
    safe_name = sanitize_filename(file_name)
    doc_id = str(uuid.uuid4())
    stored_name = f"{doc_id}{Path(safe_name).suffix.lower()}"
    file_path = UPLOAD_DIR / stored_name
    file_path.write_bytes(file_bytes)

    text, pages = extract_document_text(file_path)
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        "INSERT INTO documents (id, original_name, stored_name, extension, size_bytes, uploaded_at, file_path, status, metadata) VALUES (?, ?, ?, ?, ?, datetime('now'), ?, 'indexed', ?)",
        (
            doc_id,
            safe_name,
            stored_name,
            Path(safe_name).suffix.lower(),
            len(file_bytes),
            str(file_path),
            json.dumps({"pages": len(pages), "char_count": len(text)}),
        ),
    )

    chunks = chunk_text(text)
    for index, chunk in enumerate(chunks):
        score = _lexical_similarity(text, chunk)
        token_count = len(_tokenize(chunk))
        conn.execute(
            "INSERT INTO document_chunks (id, document_id, chunk_index, content, page_number, token_count, score) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (str(uuid.uuid4()), doc_id, index, chunk, 1 if len(pages) == 1 else index + 1, token_count, score),
        )
    conn.commit()
    conn.close()
    return {"id": doc_id, "filename": safe_name, "status": "indexed", "chunks": len(chunks)}


def delete_document(document_id: str) -> bool:
    init_db()
    conn = sqlite3.connect(DB_PATH)
    row = conn.execute("SELECT file_path FROM documents WHERE id = ?", (document_id,)).fetchone()
    if row is None:
        conn.close()
        return False
    file_path = Path(row[0])
    if file_path.exists():
        file_path.unlink(missing_ok=True)
    conn.execute("DELETE FROM document_chunks WHERE document_id = ?", (document_id,))
    conn.execute("DELETE FROM documents WHERE id = ?", (document_id,))
    conn.commit()
    conn.close()
    return True


def search_documents(query: str, limit: int = 5) -> list[dict[str, Any]]:
    init_db()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        "SELECT d.id AS document_id, d.original_name, d.file_path, c.chunk_index, c.content, c.score FROM document_chunks c JOIN documents d ON d.id = c.document_id"
    ).fetchall()
    matches = []
    for row in rows:
        score = _lexical_similarity(query, row["content"])
        if score <= 0:
            continue
        matches.append({
            "document_id": row["document_id"],
            "document_name": row["original_name"],
            "source": row["file_path"],
            "chunk_index": row["chunk_index"],
            "content": row["content"],
            "score": round(float(score), 4),
        })
    matches.sort(key=lambda item: item["score"], reverse=True)
    return matches[:limit]
